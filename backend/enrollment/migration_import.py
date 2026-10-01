"""Bounded CSV/Excel migration of a school's current operational records."""
import csv
import io
import json
import re
import hashlib
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.models import ParentStudentLink
from accounts.parent_auth import phone_value
from accounts.permissions import IsSchoolAdmin
from academics.models import AcademicSession, Term
from fees.ledger import money, post_adjustment
from fees.models import FeeCategory, FeeSchedule, StudentFinanceAccount, StudentLedgerEntry
from timetable.models import Period, TimetableEntry
from curriculum.models import AcademicStandardObjective, AcademicStandardTopic, SchoolAcademicStandard
from tenants.models import PlatformEvent, School
from .models import (ClassArm, ClassLevel, MigrationConflict, MigrationJob, MigrationMappingProfile,
                     MigrationRowRecord, MigrationStudentReference, StaffProfile, StudentProfile,
                     Subject, SubjectAssignment)
from .session_enrollment import EnrollmentPlacementError, ensure_current_enrollment
from .batch20_migration import BATCH20_DOMAINS, MigrationReview, assess_batch20, create_batch20

User = get_user_model()
DOMAINS = {
    'classes': (('class_level', 'class_arm'), ('class_level', 'class_arm')),
    'subjects': (('code', 'name'), ('code', 'name', 'class_level', 'class_levels')),
    'students': (('student_ref', 'first_name', 'last_name', 'class_level'),
                 ('student_ref', 'first_name', 'last_name', 'class_level', 'class_arm', 'email', 'dob', 'gender',
                  'state_of_origin', 'guardian_name', 'guardian_phone', 'guardian_email', 'guardian_relationship')),
    'staff': (('first_name', 'last_name', 'email'),
              ('first_name', 'last_name', 'email', 'phone', 'dob', 'gender', 'qualification',
               'specialization', 'date_employed', 'state_of_origin')),
    'parents': (('first_name', 'last_name', 'email', 'phone'),
                ('first_name', 'last_name', 'email', 'phone')),
    'parent_links': (('parent_email', 'student_ref', 'relationship'),
                     ('parent_email', 'student_ref', 'relationship')),
    'assignments': (('teacher_email', 'class_level', 'class_arm', 'subject_code'),
                    ('teacher_email', 'class_level', 'class_arm', 'subject_code')),
    'opening_balances': (('student_ref', 'amount', 'direction', 'effective_date', 'reason', 'reference'),
                         ('student_ref', 'amount', 'direction', 'effective_date', 'reason', 'reference')),
    'timetable': (('class_level', 'class_arm', 'subject_code', 'teacher_email', 'day', 'period'),
                  ('class_level', 'class_arm', 'subject_code', 'teacher_email', 'day', 'period')),
    'fee_schedules': (('class_level', 'fee_category', 'amount'),
                      ('class_level', 'fee_category', 'amount', 'due_date')),
    'standard_topics': (('standard_title', 'class_level', 'subject_code', 'term', 'position', 'title', 'recommended_week'),
                        ('standard_title', 'class_level', 'subject_code', 'term', 'position', 'title',
                         'recommended_week', 'requirement', 'description', 'source_reference', 'objectives')),
}
DOMAINS.update(BATCH20_DOMAINS)
ALIASES = {'regno': 'student_ref', 'studentnumber': 'student_ref',
           'teacheremail': 'teacher_email', 'parentemail': 'parent_email',
           'subjectcode': 'subject_code', 'classarm': 'class_arm',
           'classlevel': 'class_level', 'feecategory': 'fee_category',
           'standardtitle': 'standard_title', 'recommendedweek': 'recommended_week'}
PROTECTED = {'id', 'pk', 'school', 'schoolid', 'tenant', 'tenantid', 'userid',
             'studentid', 'staffid', 'password', 'passwordhash', 'issuperuser',
             'isstaff', 'permissions', 'groups', 'role', 'plan', 'entitlements',
             'audit', 'actor', 'admissionnumber', 'subscriptionplan',
             'platformowner', 'isactive', 'mustchangepassword', 'schoolname'}
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000


def normalized(value):
    return re.sub(r'[^a-z0-9]', '', value.casefold())


def email(value, field='email'):
    value = value.strip().lower()
    if len(value) > 254:
        raise ValueError(field, 'Email address is too long.')
    try:
        validate_email(value)
    except DjangoValidationError:
        raise ValueError(field, 'Enter a valid email address.')
    return value


def date(value, field='dob'):
    if not value:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(field, 'Use YYYY-MM-DD or DD/MM/YYYY.')


def unique(items, field, predicate, missing):
    matches = [item for item in items if predicate(item)]
    if len(matches) != 1:
        raise ValueError(field, missing if not matches else 'Reference is ambiguous; use one exact existing record.')
    return matches[0]


def level_for(school, value, *, allow_new=False):
    name = value.strip().upper().replace(' ', '')
    if name not in dict(ClassLevel.LEVEL_CHOICES):
        raise ValueError('class_level', 'Use an existing supported level such as JSS1 or SS2.')
    level = ClassLevel.objects.filter(school=school, name=name).first()
    if not level and not allow_new:
        raise ValueError('class_level', f'Create class level {name} in Classes first.')
    return name, level


def arm_for(school, level, value):
    arms = list(ClassArm.objects.filter(school=school, class_level=level))
    return unique(arms, 'class_arm',
                  lambda arm: arm.name.casefold() == value.strip().casefold() or
                  arm.full_name.casefold() == value.strip().casefold() or
                  (not value.strip() and len(arms) == 1),
                  'Choose one existing class arm in this school.')


def identify(domain, row):
    if domain == 'historical_sessions':
        return row['session'].casefold()
    if domain == 'historical_terms':
        return (row['session'].casefold(), row['term'].casefold())
    if domain == 'historical_enrollments':
        return (row['student_ref'].casefold(), row['session'].casefold(), row['enrolled_on'])
    if domain == 'historical_results':
        return (row['student_ref'].casefold(), row['session'].casefold(), row['term'].casefold(), row['subject_code'].casefold())
    if domain == 'historical_attendance':
        return (row['student_ref'].casefold(), row['session'].casefold(), row['term'].casefold(), row['date'])
    if domain == 'historical_finance':
        return (row['student_ref'].casefold(), row['reference'].casefold())
    if domain == 'opening_balances': return row['student_ref'].casefold()
    if domain == 'timetable':
        return (row['class_level'].casefold(), row['class_arm'].casefold(),
                row['day'].casefold(), row['period'].casefold())
    if domain == 'fee_schedules':
        return (row['class_level'].casefold(), row['fee_category'].casefold())
    if domain == 'standard_topics':
        return (row['standard_title'].casefold(), row['class_level'].casefold(),
                row['subject_code'].casefold(), row['term'].casefold(), row['position'])
    if domain == 'students': return row['student_ref'].casefold()
    if domain in ('staff', 'parents'): return row['email'].casefold()
    if domain == 'parent_links':
        return (row['parent_email'].casefold(), row['student_ref'].casefold())
    if domain == 'classes':
        return (row['class_level'].casefold(), row['class_arm'].casefold())
    if domain == 'subjects':
        return row['code'].casefold()
    return (row['teacher_email'].casefold(), row['class_level'].casefold(),
            row['class_arm'].casefold(), row['subject_code'].casefold())
def assess(domain, row, school, current_session=None):
    """Return (action, resolved data); perform no writes."""
    if domain in BATCH20_DOMAINS:
        return assess_batch20(domain, row, school)
    if domain == 'opening_balances':
        identity = MigrationStudentReference.objects.filter(school=school,
            reference__iexact=row['student_ref']).select_related('student').first()
        student = identity.student if identity else StudentProfile.objects.filter(school=school,
            admission_number__iexact=row['student_ref']).first()
        if not student or student.school_id != school.pk or student.user.school_id != school.pk:
            raise ValueError('student_ref', 'Match one student in this school by source reference or admission number.')
        if row['direction'].casefold() not in ('debt', 'credit'):
            raise ValueError('direction', 'Choose debt or credit.')
        try:
            amount = money(row['amount'], allow_zero=True)
        except ValueError:
            raise ValueError('amount', 'Enter a nonnegative naira amount with at most two decimal places.')
        signed = -amount if row['direction'].casefold() == 'credit' else amount
        effective = date(row['effective_date'], 'effective_date')
        from django.utils import timezone
        if not effective or effective > timezone.localdate():
            raise ValueError('effective_date', 'Use a date no later than today.')
        reason, reference = row['reason'].strip(), row['reference'].strip()
        if not reason or len(reason) > 500 or not reference or len(reference) > 100:
            raise ValueError('reason', 'Provide a reason (500 characters) and source reference (100 characters).')
        existing = StudentLedgerEntry.objects.filter(school=school, student=student, kind='opening').first()
        if existing:
            if (existing.signed_amount != signed or existing.effective_date != effective or
                    existing.reason != reason or existing.reference != reference):
                raise ValueError('student_ref', 'Opening balance already exists with different details; review manually.')
            return 'REUSE', {}
        account = StudentFinanceAccount.objects.filter(school=school, student=student).first()
        if account and account.state == 'active' and StudentLedgerEntry.objects.filter(
                school=school, student=student).exists():
            raise ValueError('student_ref', 'This account already has activity; review its opening manually.')
        return 'CREATE', {'student': student, 'signed': signed, 'effective': effective,
                          'reason': reason, 'reference': reference}
    if domain == 'timetable':
        term = Term.objects.filter(session__school=school, is_current=True).select_related('session').first()
        if not term:
            raise ValueError('term', 'Set the current academic session and term before importing timetable entries.')
        _, level = level_for(school, row['class_level'])
        arm = arm_for(school, level, row['class_arm'])
        subject = Subject.objects.filter(school=school, code__iexact=row['subject_code']).first()
        if not subject:
            raise ValueError('subject_code', 'Import this subject in this school first.')
        teacher_email = email(row['teacher_email'], 'teacher_email')
        teacher_profile = StaffProfile.objects.filter(
            school=school, user__school=school, user__role='teacher', user__is_active=True,
            employment_status='active', user__email__iexact=teacher_email
        ).select_related('user').first()
        if not teacher_profile:
            raise ValueError('teacher_email', 'Choose an active teacher in this school.')
        day_text = row['day'].strip().upper()
        aliases = {
            'MONDAY': 'MON', 'TUESDAY': 'TUE', 'WEDNESDAY': 'WED',
            'THURSDAY': 'THU', 'FRIDAY': 'FRI',
        }
        day = aliases.get(day_text, day_text[:3])
        if day not in dict(TimetableEntry.Day.choices):
            raise ValueError('day', 'Use Monday-Friday or MON-FRI.')
        period_text = row['period'].strip()
        periods = list(Period.objects.filter(school=school, is_break=False))
        period = next((
            p for p in periods
            if p.name.casefold() == period_text.casefold() or str(p.order_index) == period_text
        ), None)
        if not period:
            raise ValueError('period', 'Choose an existing non-break period by exact name or order number.')
        if not SubjectAssignment.objects.filter(
            school=school, teacher=teacher_profile, class_arm=arm, subject=subject, term=term
        ).exists():
            raise ValueError('teacher_email', 'Assign this teacher to the class and subject before importing the timetable.')
        existing = TimetableEntry.objects.filter(
            school=school, term=term, class_arm=arm, day_of_week=day, period=period
        ).first()
        if existing:
            if existing.subject_id != subject.pk or existing.teacher_id != teacher_profile.user_id:
                raise ValueError('period', 'This class already has a different timetable entry in this slot.')
            return 'REUSE', {}
        conflict = TimetableEntry.objects.filter(
            school=school, term=term, teacher=teacher_profile.user, day_of_week=day, period=period
        ).first()
        if conflict:
            raise ValueError('teacher_email', f'Teacher is already scheduled for {conflict.class_arm} in this slot.')
        return 'CREATE', {
            'term': term, 'arm': arm, 'subject': subject, 'teacher': teacher_profile.user,
            'day': day, 'period': period,
        }

    if domain == 'fee_schedules':
        term = Term.objects.filter(session__school=school, is_current=True).first()
        if not term:
            raise ValueError('term', 'Set the current academic term before importing fee schedules.')
        _, level = level_for(school, row['class_level'])
        category = FeeCategory.objects.filter(
            school=school, name__iexact=row['fee_category'].strip()
        ).first()
        if not category:
            raise ValueError('fee_category', 'Create this fee category before importing schedules.')
        try:
            amount = Decimal(row['amount'])
            if not amount.is_finite() or amount <= 0 or amount != amount.quantize(Decimal('0.01')):
                raise ValueError()
        except Exception:
            raise ValueError('amount', 'Enter a positive fee amount with at most two decimal places.')
        due = date(row.get('due_date', ''), 'due_date')
        existing = FeeSchedule.objects.filter(
            school=school, term=term, class_level=level, fee_category=category
        ).first()
        if existing:
            if existing.amount != amount or existing.due_date != due:
                raise ValueError('amount', 'A fee schedule already exists with different details; review it manually.')
            return 'REUSE', {}
        return 'CREATE', {
            'term': term, 'level': level, 'category': category,
            'amount': amount, 'due_date': due,
        }

    if domain == 'standard_topics':
        _, level = level_for(school, row['class_level'])
        subject = Subject.objects.filter(school=school, code__iexact=row['subject_code']).first()
        if not subject:
            raise ValueError('subject_code', 'Import this subject in this school first.')
        standard = SchoolAcademicStandard.objects.filter(
            school=school, class_level=level, subject=subject,
            title__iexact=row['standard_title'].strip(),
            status=SchoolAcademicStandard.Status.DRAFT,
        ).order_by('-revision').first()
        if not standard:
            raise ValueError('standard_title', 'Choose an existing draft academic standard with this title/class/subject.')
        term_name = row['term'].strip().casefold()
        if term_name not in ('first', 'second', 'third'):
            raise ValueError('term', 'Use first, second or third.')
        try:
            position = int(row['position'])
            week = int(row['recommended_week'])
            if position < 1 or position > 200 or week < 1 or week > 52:
                raise ValueError()
        except Exception:
            raise ValueError('position', 'Position must be 1-200 and recommended_week must be 1-52.')
        requirement = row.get('requirement', 'required').strip().casefold() or 'required'
        if requirement not in dict(AcademicStandardTopic.Requirement.choices):
            raise ValueError('requirement', 'Use required or enrichment.')
        title = row['title'].strip()
        if not title or len(title) > 180:
            raise ValueError('title', 'Provide a topic title of at most 180 characters.')
        objectives = [item.strip() for item in row.get('objectives', '').split(';') if item.strip()]
        if len(objectives) > 30 or any(len(item) > 300 for item in objectives):
            raise ValueError('objectives', 'Use up to 30 semicolon-separated objectives, each at most 300 characters.')
        if len({item.casefold() for item in objectives}) != len(objectives):
            raise ValueError('objectives', 'Duplicate objectives are not allowed.')
        existing = AcademicStandardTopic.objects.filter(
            standard=standard, term=term_name, position=position
        ).first()
        if existing:
            if existing.title.casefold() != title.casefold():
                raise ValueError('position', 'This standard/term position already has a different topic.')
            return 'REUSE', {}
        return 'CREATE', {
            'standard': standard, 'term': term_name, 'position': position, 'week': week,
            'title': title, 'requirement': requirement, 'objectives': objectives,
        }

    if domain == 'classes':
        name, level = level_for(school, row['class_level'], allow_new=True)
        arm_name = row['class_arm'].strip()
        if len(arm_name) > 10: raise ValueError('class_arm', 'Use at most 10 characters.')
        existing = ClassArm.objects.filter(school=school, class_level=level) if level else []
        arm = next((a for a in existing if a.name.casefold() == arm_name.casefold()), None)
        return ('REUSE' if arm else 'CREATE', {'level_name': name, 'arm_name': arm_name})
    if domain == 'subjects':
        code, name = row['code'].upper(), row['name']
        if len(code) > 10 or len(name) > 100: raise ValueError('code', 'Check subject code (10) and name (100) lengths.')
        if row.get('class_level') and row.get('class_levels'):
            raise ValueError('class_levels', 'Use class_level or class_levels, not both.')
        names = row.get('class_levels', '').split(';') if row.get('class_levels') else [row.get('class_level')] if row.get('class_level') else []
        levels = [level_for(school, item)[1] for item in names]
        if len(levels) != len({item.pk for item in levels}):
            raise ValueError('class_levels', 'List each class level once.')
        subject = Subject.objects.filter(school=school, code__iexact=code).first()
        if subject and subject.name.casefold() != name.casefold():
            raise ValueError('code', 'This subject code already has a different name; review it manually.')
        if subject and levels and set(subject.class_levels.values_list('pk', flat=True)) != {item.pk for item in levels}:
            raise ValueError('class_level', 'Existing subject has different class coverage; review it manually.')
        return ('REUSE' if subject else 'CREATE', {'code': code, 'name': name, 'levels': levels})
    if domain == 'students':
        ref = row['student_ref']
        if len(ref) > 80: raise ValueError('student_ref', 'Use at most 80 characters.')
        _, level = level_for(school, row['class_level'])
        arm = arm_for(school, level, row.get('class_arm', ''))
        dob = date(row.get('dob', ''))
        gender = row.get('gender', '').casefold()
        if gender not in ('', 'male', 'female', 'm', 'f'):
            raise ValueError('gender', 'Use male or female.')
        student_email = email(row['email']) if row.get('email') else None
        if row.get('guardian_email'): email(row['guardian_email'], 'guardian_email')
        for field, limit in (('state_of_origin', 50), ('guardian_name', 150), ('guardian_phone', 20)):
            if len(row.get(field, '')) > limit: raise ValueError(field, f'Use at most {limit} characters.')
        relationship = row.get('guardian_relationship', '').casefold()
        if relationship and relationship not in dict(StudentProfile.RELATIONSHIP_CHOICES):
            raise ValueError('guardian_relationship', 'Choose a supported guardian relationship.')
        identity = MigrationStudentReference.objects.filter(school=school, reference__iexact=ref).select_related('student__user').first()
        if identity:
            student = identity.student
            if (student.school_id != school.pk or student.user.school_id != school.pk or
                student.user.first_name.casefold() != row['first_name'].casefold() or
                student.user.last_name.casefold() != row['last_name'].casefold() or
                student.current_class_id != arm.pk or student.dob != dob or
                (student_email and student.user.email != student_email)):
                raise ValueError('student_ref', 'Reference is already linked to different student details; review manually.')
            return 'REUSE', {}
        if student_email and User.objects.filter(email__iexact=student_email).exists():
            raise ValueError('email', 'This email cannot be used; review the existing account.')
        if StudentProfile.objects.filter(school=school, user__first_name__iexact=row['first_name'],
                user__last_name__iexact=row['last_name'], dob=dob, current_class=arm).exists():
            raise ValueError('student_ref', 'A matching student already exists; review before linking a source reference.')
        if arm and not current_session:
            raise ValueError(
                'class_arm',
                "Set the school's current academic session before importing students into classes.",
            )
        return 'CREATE', {
            'arm': arm,
            'dob': dob,
            'gender': {'m': 'male', 'f': 'female'}.get(gender, gender),
            'email': student_email,
            'current_session': current_session,
        }
    if domain == 'staff':
        if len(row.get('phone', '')) > 20: raise ValueError('phone', 'Use at most 20 characters.')
        if len(row.get('specialization', '')) > 150 or len(row.get('state_of_origin', '')) > 50:
            raise ValueError('specialization', 'Check specialization and state of origin lengths.')
        gender = row.get('gender', '').casefold()
        qualification = row.get('qualification', '').casefold()
        if gender and gender not in dict(StaffProfile.GENDER_CHOICES):
            raise ValueError('gender', 'Use male or female.')
        if qualification and qualification not in dict(StaffProfile.QUALIFICATION_CHOICES):
            raise ValueError('qualification', 'Choose a supported qualification code.')
        dob = date(row.get('dob', ''))
        employed = date(row.get('date_employed', ''), 'date_employed')
        staff_email = email(row['email'])
        user = User.objects.filter(email__iexact=staff_email).first()
        if user:
            if (user.school_id != school.pk or user.role != 'teacher' or
                user.first_name.casefold() != row['first_name'].casefold() or
                user.last_name.casefold() != row['last_name'].casefold() or
                not StaffProfile.objects.filter(school=school, user=user).exists()):
                raise ValueError('email', 'This email belongs to a different account; review manually.')
            return 'REUSE', {}
        return 'CREATE', {'email': staff_email, 'dob': dob, 'employed': employed,
                          'gender': gender, 'qualification': qualification}
    if domain == 'parents':
        parent_email = email(row['email'])
        phone = phone_value(row['phone'])
        if not phone: raise ValueError('phone', 'Enter a parent login phone number with 10 to 15 digits.')
        user = User.objects.filter(email__iexact=parent_email).first()
        if user:
            if (user.school_id != school.pk or user.role != 'parent' or
                user.phone_number != phone or not user.is_active or
                user.first_name.casefold() != row['first_name'].casefold() or
                user.last_name.casefold() != row['last_name'].casefold()):
                raise ValueError('email', 'This parent email has different account details; review manually.')
            return 'REUSE', {}
        if User.objects.filter(school=school, role='parent', phone_number=phone).exists():
            raise ValueError('phone', 'This phone belongs to another parent account; use its registered email.')
        return 'CREATE', {'email': parent_email, 'phone': phone}
    if domain == 'parent_links':
        parent = User.objects.filter(school=school, role='parent', email__iexact=email(row['parent_email'], 'parent_email'), is_active=True).first()
        if not parent: raise ValueError('parent_email', 'Import this parent in this school first.')
        identity = MigrationStudentReference.objects.filter(school=school, reference__iexact=row['student_ref']).select_related('student__user').first()
        student = identity.student if identity else StudentProfile.objects.filter(school=school, admission_number__iexact=row['student_ref']).select_related('user').first()
        if not student or student.school_id != school.pk or student.user.school_id != school.pk:
            raise ValueError('student_ref', 'Import this student reference in this school first.')
        if row['relationship'].casefold() not in dict(ParentStudentLink.RELATIONSHIP_CHOICES):
            raise ValueError('relationship', 'Choose father, mother or guardian.')
        link = ParentStudentLink.objects.filter(school=school, parent=parent, student=student).first()
        if link and link.relationship != row['relationship'].casefold():
            raise ValueError('relationship', 'This link exists with another relationship; review manually.')
        return ('REUSE' if link else 'CREATE', {'parent': parent, 'student': student})
    teacher_email = email(row['teacher_email'], 'teacher_email')
    teacher = StaffProfile.objects.filter(school=school, user__school=school, user__role='teacher',
                                          user__email__iexact=teacher_email).first()
    if not teacher: raise ValueError('teacher_email', 'Import this teacher in this school first.')
    _, level = level_for(school, row['class_level'])
    arm = arm_for(school, level, row['class_arm'])
    subject = Subject.objects.filter(school=school, code__iexact=row['subject_code']).first()
    if not subject: raise ValueError('subject_code', 'Import this subject in this school first.')
    if subject.class_levels.exists() and not subject.class_levels.filter(pk=level.pk).exists():
        raise ValueError('subject_code', 'This subject is not offered at the selected class level.')
    term = Term.objects.filter(session__school=school, is_current=True).select_related('session').first()
    if not term: raise ValueError('term', 'Set the current academic session and term before importing assignments.')
    existing = SubjectAssignment.objects.filter(school=school, subject=subject, class_arm=arm, term=term).first()
    if existing and existing.teacher_id != teacher.pk:
        raise ValueError('teacher_email', 'This subject/class/term already belongs to another teacher; review manually.')
    return ('REUSE' if existing else 'CREATE', {'teacher': teacher, 'arm': arm, 'subject': subject, 'term': term})


def create(domain, row, school, data):
    if domain in BATCH20_DOMAINS:
        return create_batch20(domain, row, school, data, actor=data.get('actor'))
    if domain == 'opening_balances':
        try:
            post_adjustment(school, data['student'], data['actor'], kind='opening',
                amount=data['signed'], reason=data['reason'], reference=data['reference'],
                effective_date=data['effective'], key=f'opening-import:{school.pk}:{data["student"].pk}')
        except ValueError as exc:
            raise ValueError('student_ref', str(exc)) from exc
    elif domain == 'timetable':
        TimetableEntry.objects.create(
            school=school, term=data['term'], class_arm=data['arm'], subject=data['subject'],
            teacher=data['teacher'], day_of_week=data['day'], period=data['period']
        )
    elif domain == 'fee_schedules':
        FeeSchedule.objects.create(
            school=school, term=data['term'], class_level=data['level'], fee_category=data['category'],
            amount=data['amount'], due_date=data['due_date']
        )
    elif domain == 'standard_topics':
        topic = AcademicStandardTopic.objects.create(
            standard=data['standard'], term=data['term'], title=data['title'],
            description=row.get('description', '')[:500], position=data['position'],
            recommended_week=data['week'], requirement=data['requirement'],
            source_reference=row.get('source_reference', '')[:180]
        )
        AcademicStandardObjective.objects.bulk_create([
            AcademicStandardObjective(topic=topic, text=text, position=index)
            for index, text in enumerate(data['objectives'], 1)
        ])
    elif domain == 'classes':
        level, _ = ClassLevel.objects.get_or_create(school=school, name=data['level_name'])
        ClassArm.objects.create(school=school, class_level=level, name=data['arm_name'])
    elif domain == 'subjects':
        subject = Subject.objects.create(school=school, code=data['code'], name=data['name'])
        if data['levels']: subject.class_levels.add(*data['levels'])
    elif domain == 'students':
        user = User.objects.create_user(email=data['email'], password=None,
            first_name=row['first_name'], last_name=row['last_name'], role='student',
            school=school, must_change_password=True)
        student = StudentProfile.objects.create(
            user=user,
            school=school,
            current_class=None,
            dob=data['dob'],
            gender=data['gender'],
            state_of_origin=row.get('state_of_origin', ''),
            guardian_name=row.get('guardian_name', ''),
            guardian_phone=row.get('guardian_phone', ''),
            guardian_email=row.get('guardian_email', '').lower(),
            guardian_relationship=row.get('guardian_relationship', '').lower(),
        )
        if data['arm']:
            ensure_current_enrollment(
                school=school,
                student=student,
                class_arm=data['arm'],
                actor=data.get('actor'),
                entry_reason='admission',
                current_session=data.get('current_session'),
            )
        user.set_password(student.admission_number)
        user.save(update_fields=['password'])
        MigrationStudentReference.objects.create(school=school, reference=row['student_ref'], student=student)
    elif domain == 'staff':
        user = User.objects.create_user(email=data['email'], password=None,
            first_name=row['first_name'], last_name=row['last_name'], role='teacher',
            school=school, must_change_password=True)
        staff = StaffProfile.objects.create(user=user, school=school, phone=row.get('phone', ''),
            dob=data['dob'], date_employed=data['employed'], gender=data['gender'],
            qualification=data['qualification'], specialization=row.get('specialization', ''),
            state_of_origin=row.get('state_of_origin', ''))
        user.set_password(staff.staff_id)
        user.save(update_fields=['password'])
    elif domain == 'parents':
        User.objects.create_user(email=data['email'], password=None, first_name=row['first_name'],
            last_name=row['last_name'], role='parent', school=school, phone_number=data['phone'],
            must_change_password=False)
    elif domain == 'parent_links':
        ParentStudentLink.objects.create(school=school, parent=data['parent'],
            student=data['student'], relationship=row['relationship'].casefold())
    else:
        SubjectAssignment.objects.create(school=school, teacher=data['teacher'],
            subject=data['subject'], class_arm=data['arm'], term=data['term'], session=data['term'].session)


def _tabular_value(cell):
    value = cell.value if hasattr(cell, 'value') else cell
    if value is None:
        return ''
    if hasattr(cell, 'data_type') and cell.data_type == 'f':
        return '=' + str(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _upload_fingerprint(upload):
    try:
        upload.seek(0)
        blob = upload.read(MAX_BYTES + 1)
        upload.seek(0)
    except Exception:
        raise ValidationError({'error': 'Could not read the uploaded file.'})
    if len(blob) > MAX_BYTES:
        raise ValidationError({'error': 'Spreadsheet is larger than 2 MB.'})
    return hashlib.sha256(blob).hexdigest()


def _job_for_upload(request, domain, *, status='inspected', total_rows=0, file_format='', mapping=None):
    upload = request.FILES.get('file')
    if not upload:
        raise ValidationError({'error': 'Choose a CSV or Excel .xlsx file.'})
    fingerprint = _upload_fingerprint(upload)
    defaults = {
        'original_filename': (upload.name or 'upload')[:255],
        'file_format': file_format,
        'status': status,
        'total_rows': total_rows,
        'created_by': request.user,
    }
    if mapping is not None:
        defaults['mapping_snapshot'] = mapping
    job, _ = MigrationJob.objects.get_or_create(
        school=request.tenant,
        domain=domain,
        file_fingerprint=fingerprint,
        defaults=defaults,
    )
    terminal = job.status in ('completed', 'completed_with_errors')
    if terminal:
        return job, fingerprint
    changed = []
    for field, value in defaults.items():
        if field == 'created_by':
            continue
        if value not in ('', None) and getattr(job, field) != value:
            setattr(job, field, value)
            changed.append(field)
    if changed:
        job.save(update_fields=changed + ['updated_at'])
    return job, fingerprint


def _read_tabular_upload(upload):
    name = (upload.name or '').lower()
    if not name.endswith(('.csv', '.xlsx')) or upload.size > MAX_BYTES:
        raise ValidationError({'error': 'Upload a CSV or Excel .xlsx file of at most 2 MB.'})
    blob = upload.read(MAX_BYTES + 1)
    if len(blob) > MAX_BYTES:
        raise ValidationError({'error': 'Spreadsheet is larger than 2 MB.'})

    if name.endswith('.csv'):
        try:
            content = blob.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValidationError({'error': 'Save CSV files as UTF-8, or upload an Excel .xlsx file.'})
        if '\x00' in content:
            raise ValidationError({'error': 'CSV contains invalid characters.'})
        try:
            reader = csv.DictReader(io.StringIO(content, newline=''), strict=True)
            headers = reader.fieldnames or []
            indexed_rows = [(number, raw) for number, raw in enumerate(reader, 2)]
        except csv.Error:
            raise ValidationError({'error': 'CSV structure is malformed.'})
        if len(indexed_rows) > MAX_ROWS:
            raise ValidationError({'error': 'Import at most 2,000 rows per file.'})
        return headers, indexed_rows, 'csv'

    try:
        workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=False)
        sheet = workbook.active
        iterator = sheet.iter_rows()
        header_cells = next(iterator, None)
        headers = [_tabular_value(cell).strip() for cell in (header_cells or [])]
        indexed_rows = []
        for worksheet_row, cells in enumerate(iterator, 2):
            values = [_tabular_value(cell) for cell in cells[:len(headers)]]
            if not any(value.strip() for value in values):
                continue
            if len(indexed_rows) >= MAX_ROWS:
                raise ValidationError({'error': 'Import at most 2,000 rows per file.'})
            indexed_rows.append((worksheet_row, dict(zip(headers, values))))
        workbook.close()
    except ValidationError:
        raise
    except Exception:
        raise ValidationError({'error': 'Excel workbook could not be read. Upload a valid .xlsx file.'})
    return headers, indexed_rows, 'xlsx'


def _validate_headers(headers):
    if not headers or any(not h.strip() for h in headers) or len(headers) != len(set(headers)) or len(headers) != len(set(map(normalized, headers))):
        raise ValidationError({'error': 'Spreadsheet headers must be present and unique.'})
    if any(normalized(h) in PROTECTED or any(word in normalized(h) for word in ('password', 'secret', 'token')) for h in headers):
        raise ValidationError({'error': 'Remove system-controlled columns such as IDs, roles, passwords or school.'})


def _suggest_mapping(domain, headers):
    allowed = set(DOMAINS[domain][1])
    mapping = {}
    for header in headers:
        normalized_header = normalized(header)
        candidate = next((field for field in allowed if normalized(field) == normalized_header), None)
        candidate = candidate or ALIASES.get(normalized_header)
        if candidate in allowed and candidate not in mapping.values():
            mapping[header] = candidate
    return mapping


def inspect_upload(request, domain):
    if domain not in DOMAINS:
        raise ValidationError({'error': 'Unsupported migration type.'})
    upload = request.FILES.get('file')
    if not upload:
        raise ValidationError({'error': 'Choose a CSV or Excel .xlsx file.'})
    headers, indexed_rows, file_format = _read_tabular_upload(upload)
    _validate_headers(headers)
    suggested = _suggest_mapping(domain, headers)
    matched_profile = None
    profiles = MigrationMappingProfile.objects.filter(
        school=request.tenant, domain=domain
    ).order_by('-last_used_at', '-updated_at', '-id')
    for profile in profiles:
        profile_mapping = profile.mappings if isinstance(profile.mappings, dict) else {}
        if profile_mapping and set(profile_mapping).issubset(set(headers)):
            values = [value for value in profile_mapping.values() if value]
            if len(values) == len(set(values)) and all(value in DOMAINS[domain][1] for value in values):
                suggested.update(profile_mapping)
                matched_profile = profile
                break
    job, fingerprint = _job_for_upload(
        request, domain, status='inspected', total_rows=len(indexed_rows),
        file_format=file_format, mapping=suggested,
    )
    return {
        'headers': headers,
        'rows': [raw for _, raw in indexed_rows],
        'row_numbers': [number for number, _ in indexed_rows],
        'total_rows': len(indexed_rows),
        'suggested_mapping': suggested,
        'format': file_format,
        'job_id': job.pk,
        'file_fingerprint': fingerprint,
        'mapping_profile': ({
            'id': matched_profile.pk,
            'name': matched_profile.name,
            'source_system': matched_profile.source_system,
        } if matched_profile else None),
    }


def parse_upload(request, domain):
    if domain not in DOMAINS:
        raise ValidationError({'error': 'Unsupported migration type.'})
    upload = request.FILES.get('file')
    if not upload:
        raise ValidationError({'error': 'Choose a CSV or Excel .xlsx file.'})
    headers, indexed_rows, _ = _read_tabular_upload(upload)
    _validate_headers(headers)
    try:
        supplied = json.loads(request.data.get('mapping', '{}'))
    except (TypeError, ValueError):
        raise ValidationError({'error': 'Column mapping must be valid JSON.'})
    if not isinstance(supplied, dict) or any(not isinstance(source, str) or source not in headers for source in supplied):
        raise ValidationError({'error': 'Column mapping refers to a missing source column.'})
    allowed = set(DOMAINS[domain][1])
    chosen = [dest for dest in supplied.values() if dest not in ('', None)]
    if any(not isinstance(dest, str) or dest not in allowed for dest in chosen) or len(set(chosen)) != len(chosen):
        raise ValidationError({'error': 'Mapping has an unsupported or duplicate destination.'})
    mapping = {source: dest for source, dest in supplied.items() if dest not in ('', None)}
    for source, dest in _suggest_mapping(domain, headers).items():
        if source not in supplied and dest not in mapping.values():
            mapping[source] = dest
    if set(DOMAINS[domain][0]) - set(mapping.values()):
        raise ValidationError({'error': 'Map required columns: ' + ', '.join(sorted(set(DOMAINS[domain][0]) - set(mapping.values())))})
    rows = []
    for number, raw in indexed_rows:
        if None in raw or any(value is None for value in raw.values()):
            rows.append((number, None, ('file', 'Column count does not match the header.')))
            continue
        row = {dest: str(raw[source]).strip() for source, dest in mapping.items()}
        if any(value.lstrip().startswith(('=', '@')) for value in row.values()):
            rows.append((number, None, ('file', 'Formula-like content is not accepted in migration fields.')))
        else:
            rows.append((number, row, None))
    job, fingerprint = _job_for_upload(
        request, domain, total_rows=len(rows), mapping=mapping,
    )
    return rows, [h for h in headers if h not in mapping], mapping, job, fingerprint


class MigrationCentre(APIView):
    permission_classes = [IsSchoolAdmin]
    parser_classes = [MultiPartParser]

    def get(self, request, domain=None):
        if domain:
            if domain not in DOMAINS: raise ValidationError({'error': 'Unsupported migration type.'})
            from django.http import HttpResponse
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(DOMAINS[domain][1])
            response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = f'attachment; filename="paideia_{domain}_template.csv"'
            return response
        return Response({'domains': [
            {'key': key, 'required': required, 'columns': columns}
            for key, (required, columns) in DOMAINS.items()]})

    def post(self, request, domain, operation):
        if operation not in ('inspect', 'validate', 'import'):
            raise ValidationError({'error': 'Choose inspect, validate or import.'})
        if operation == 'inspect':
            return Response(inspect_upload(request, domain))
        rows, ignored, mapping, job, fingerprint = parse_upload(request, domain)
        job_was_terminal = job.status in ('completed', 'completed_with_errors')
        school = request.tenant
        current_session = (
            AcademicSession.objects.filter(school=school, is_current=True).first()
            if domain == 'students'
            else None
        )
        results, seen, seen_emails, seen_people = [], set(), set(), set()
        seen_phones, seen_assignment_slots = set(), set()
        seen_timetable_teacher_slots = set()
        counts = {'CREATE': 0, 'REUSE': 0, 'REVIEW': 0, 'REJECT': 0}

        def process():
            for number, row, parse_error in rows:
                target = None
                if parse_error:
                    outcome = {'row': number, 'action': 'REJECT', 'field': parse_error[0], 'reason': parse_error[1]}
                else:
                    try:
                        missing = [field for field in DOMAINS[domain][0] if not row.get(field)]
                        if missing: raise ValueError(missing[0], 'Required value is empty.')
                        for field in ('first_name', 'last_name'):
                            if len(row.get(field, '')) > 150:
                                raise ValueError(field, 'Use at most 150 characters.')
                        key = identify(domain, row)
                        if key in seen: raise ValueError('file', 'Duplicate reference in this file.')
                        seen.add(key)
                        action, data = assess(
                            domain,
                            row,
                            school,
                            current_session=current_session,
                        )
                        if domain == 'parents' and action == 'CREATE':
                            if data['phone'] in seen_phones:
                                raise ValueError('phone', 'Duplicate parent phone in this file.')
                            seen_phones.add(data['phone'])
                        if domain == 'assignments' and action == 'CREATE':
                            slot = (data['arm'].pk, data['subject'].pk, data['term'].pk)
                            if slot in seen_assignment_slots:
                                raise ValueError('subject_code', 'Another teacher in this file already has this class and subject.')
                            seen_assignment_slots.add(slot)
                        if domain == 'timetable' and action == 'CREATE':
                            teacher_slot = (data['teacher'].pk, data['day'], data['period'].pk, data['term'].pk)
                            if teacher_slot in seen_timetable_teacher_slots:
                                raise ValueError('teacher_email', 'This teacher is assigned to more than one class in the same timetable slot in this file.')
                            seen_timetable_teacher_slots.add(teacher_slot)
                        if domain == 'students' and action == 'CREATE':
                            student_email = data['email']
                            person = (row['first_name'].casefold(), row['last_name'].casefold(),
                                      data['dob'], data['arm'].pk)
                            if student_email and student_email in seen_emails:
                                raise ValueError('email', 'Duplicate email in this file.')
                            if person in seen_people:
                                raise ValueError('student_ref', 'Duplicate student details in this file.')
                            if student_email: seen_emails.add(student_email)
                            seen_people.add(person)
                        target = None
                        if operation == 'import' and action == 'CREATE':
                            with transaction.atomic():
                                if domain in ('opening_balances', 'students') or domain in BATCH20_DOMAINS:
                                    data['actor'] = request.user
                                target = create(domain, row, school, data)
                        outcome = {'row': number, 'action': action}
                    except MigrationReview as exc:
                        MigrationConflict.objects.update_or_create(
                            job=job,
                            row_number=number,
                            conflict_type=exc.conflict_type,
                            defaults={
                                'domain': domain,
                                'source_identity': str(identify(domain, row))[:255] if row else '',
                                'source_payload': row or {},
                                'candidate_matches': exc.candidates,
                                'status': 'open',
                            },
                        )
                        outcome = {'row': number, 'action': 'REVIEW', 'field': exc.field, 'reason': exc.reason}
                    except ValueError as exc:
                        outcome = {'row': number, 'action': 'REJECT', 'field': exc.args[0], 'reason': exc.args[1]}
                    except Exception:
                        if operation == 'validate': raise
                        outcome = {'row': number, 'action': 'REJECT', 'field': 'file',
                                   'reason': 'This row could not be saved. Check its values and retry it separately.'}
                counts[outcome['action']] += 1
                results.append(outcome)
                if operation == 'import' and not job_was_terminal:
                    source_identity = ''
                    if row:
                        try:
                            source_identity = str(identify(domain, row))[:255]
                        except Exception:
                            source_identity = ''
                    target_obj = locals().get('target')
                    MigrationRowRecord.objects.update_or_create(
                        job=job,
                        row_number=number,
                        defaults={
                            'action': outcome['action'],
                            'source_identity': source_identity,
                            'target_model': (
                                target_obj._meta.label_lower
                                if target_obj is not None and hasattr(target_obj, '_meta')
                                else ''
                            ),
                            'target_pk': str(getattr(target_obj, 'pk', '') or ''),
                        },
                    )
                    if number > job.last_processed_row:
                        job.last_processed_row = number

        if operation == 'import':
            if not job_was_terminal:
                job.status = 'importing'
                job.started_at = timezone.now()
                job.mapping_snapshot = mapping
                job.save(update_fields=['status', 'started_at', 'mapping_snapshot', 'updated_at'])
            process()
            if not job_was_terminal:
                job.status = 'completed_with_errors' if (counts['REJECT'] or counts['REVIEW']) else 'completed'
                job.completed_at = timezone.now()
        else:
            process()
            if not job_was_terminal:
                job.status = 'validated'
                job.validated_at = timezone.now()
        if not job_was_terminal:
            job.total_rows = len(rows)
            job.create_count = counts['CREATE']
            job.reuse_count = counts['REUSE']
            job.review_count = counts['REVIEW']
            job.reject_count = counts['REJECT']
            job.mapping_snapshot = mapping
            job.save(update_fields=[
                'status', 'total_rows', 'create_count', 'reuse_count', 'review_count',
                'reject_count', 'last_processed_row', 'mapping_snapshot', 'validated_at',
                'started_at', 'completed_at', 'updated_at',
            ])
        if operation == 'import' and not job_was_terminal and (counts['CREATE'] or counts['REVIEW'] or counts['REJECT']):
            PlatformEvent.objects.create(
                actor=request.user, actor_email=request.user.email,
                action='school.migration_completed', target=str(job.pk),
                details={
                    'school_id': school.pk, 'domain': domain, 'counts': counts,
                    'migration_job_id': job.pk, 'file_fingerprint': fingerprint,
                },
            )
        return Response({'domain': domain, 'mode': operation, 'total_rows': len(rows),
                         'counts': counts, 'rows': results,
                         'warnings': [f'Ignored column: {header}' for header in ignored],
                         'mapping': mapping, 'job_id': job.pk,
                         'file_fingerprint': fingerprint, 'status': job.status,
                         'idempotent_replay': bool(operation == 'import' and job_was_terminal)})


class MigrationJobList(APIView):
    permission_classes = [IsSchoolAdmin]

    def get(self, request):
        rows = MigrationJob.objects.filter(school=request.tenant).select_related('created_by')[:100]
        return Response([{
            'id': row.pk, 'domain': row.domain, 'source_name': row.source_name,
            'original_filename': row.original_filename, 'file_fingerprint': row.file_fingerprint,
            'status': row.status, 'total_rows': row.total_rows,
            'last_processed_row': row.last_processed_row,
            'create_count': row.create_count, 'reuse_count': row.reuse_count,
            'review_count': row.review_count, 'reject_count': row.reject_count,
            'created_at': row.created_at, 'validated_at': row.validated_at,
            'completed_at': row.completed_at,
            'created_by': row.created_by.email if row.created_by else None,
        } for row in rows])


class MigrationJobReport(APIView):
    permission_classes = [IsSchoolAdmin]

    def get(self, request, pk):
        from django.http import HttpResponse
        job = MigrationJob.objects.filter(pk=pk, school=request.tenant).first()
        if not job:
            raise ValidationError({'error': 'Migration job not found.'})
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['row_number', 'action', 'source_identity', 'target_model', 'target_pk'])
        for row in job.row_records.all().order_by('row_number'):
            writer.writerow([
                row.row_number, row.action, row.source_identity,
                row.target_model, row.target_pk,
            ])
        response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="paideia_migration_job_{job.pk}.csv"'
        return response


class MigrationMappingProfiles(APIView):
    permission_classes = [IsSchoolAdmin]

    def get(self, request):
        qs = MigrationMappingProfile.objects.filter(school=request.tenant)
        domain = request.query_params.get('domain')
        if domain:
            qs = qs.filter(domain=domain)
        return Response([{
            'id': row.pk, 'domain': row.domain, 'name': row.name,
            'source_system': row.source_system, 'mappings': row.mappings,
            'last_used_at': row.last_used_at,
        } for row in qs])

    def post(self, request):
        domain = str(request.data.get('domain') or '').strip()
        name = str(request.data.get('name') or '').strip()
        mappings = request.data.get('mappings')
        if domain not in DOMAINS or not name or not isinstance(mappings, dict):
            raise ValidationError({'error': 'Provide a supported domain, profile name and mappings object.'})
        allowed = set(DOMAINS[domain][1])
        if any(key not in allowed for key in mappings.values() if key):
            raise ValidationError({'error': 'Mapping profile contains an unsupported destination field.'})
        profile, _ = MigrationMappingProfile.objects.update_or_create(
            school=request.tenant, domain=domain, name=name,
            defaults={
                'source_system': str(request.data.get('source_system') or '')[:120],
                'mappings': mappings, 'created_by': request.user,
                'last_used_at': timezone.now(),
            },
        )
        return Response({'id': profile.pk, 'domain': profile.domain, 'name': profile.name,
                         'source_system': profile.source_system, 'mappings': profile.mappings}, status=201)


class MigrationConflictList(APIView):
    permission_classes = [IsSchoolAdmin]

    def get(self, request):
        qs = MigrationConflict.objects.filter(job__school=request.tenant).select_related('job', 'resolved_by')
        status_filter = request.query_params.get('status')
        if status_filter:
            qs = qs.filter(status=status_filter)
        return Response([{
            'id': row.pk, 'job_id': row.job_id, 'domain': row.domain,
            'row_number': row.row_number, 'source_identity': row.source_identity,
            'conflict_type': row.conflict_type, 'source_payload': row.source_payload,
            'candidate_matches': row.candidate_matches, 'status': row.status,
            'resolution': row.resolution, 'resolved_at': row.resolved_at,
            'resolved_by': row.resolved_by.email if row.resolved_by else None,
        } for row in qs[:200]])

    def patch(self, request):
        conflict = MigrationConflict.objects.filter(
            pk=request.data.get('id'), job__school=request.tenant
        ).first()
        if not conflict:
            raise ValidationError({'error': 'Migration conflict not found.'})
        action = request.data.get('action')
        if action not in ('resolved', 'ignored'):
            raise ValidationError({'error': 'Choose resolved or ignored.'})
        resolution = request.data.get('resolution') or {}
        if not isinstance(resolution, dict):
            raise ValidationError({'error': 'Resolution must be an object.'})
        conflict.status = action
        conflict.resolution = resolution
        conflict.resolved_by = request.user
        conflict.resolved_at = timezone.now()
        conflict.save(update_fields=['status', 'resolution', 'resolved_by', 'resolved_at'])
        PlatformEvent.objects.create(
            actor=request.user, actor_email=request.user.email,
            action='school.migration_conflict_resolved', target=str(conflict.pk),
            details={'school_id': request.tenant.pk, 'job_id': conflict.job_id,
                     'status': action, 'resolution': resolution},
        )
        return Response({'id': conflict.pk, 'status': conflict.status, 'resolution': conflict.resolution})
